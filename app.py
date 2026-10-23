import os
import re
import json
import time
import base64
import secrets
import threading
import tempfile
import html
import hashlib
from collections import Counter
from datetime import datetime, date, timedelta
from pathlib import Path

import requests
from flask import Flask, request, jsonify
from zoneinfo import ZoneInfo
from fleet_bot import install as install_fleet
from report_menu import install as install_report_menu
import storage_v5_1 as db

app = Flask(__name__)

# ============================================================
# SGSB BAGGING BOT V5.1.1
# - Permanent Outlet Master (Supabase)
# - Persistent bagging metadata and staff sessions (Supabase)
# - Multiple/parallel Record Cards
# - 1 Seal = 1 Photo + 1 Video
# - Telegram album archive (photo + video grouped together)
# - Late/backdated upload and resume incomplete record
# - Detailed on-demand reports by date range / outlet
# - Report output: Telegram summary/detail, PDF, Excel, TXT, PNG chart
# ============================================================

# ============================================================
# ENVIRONMENT
# ============================================================

BOT_TOKEN = os.environ.get("BOT_TOKEN", "")
ARCHIVE_CHAT_ID = os.environ.get("ARCHIVE_CHAT_ID", "")
WEBHOOK_SECRET = os.environ.get("WEBHOOK_SECRET", "")
ADMIN_PASSCODE = os.environ.get("ADMIN_PASSCODE", "")

API = f"https://api.telegram.org/bot{BOT_TOKEN}"
TZ = ZoneInfo("Asia/Kuala_Lumpur")
REPORT_DIR = Path(tempfile.gettempdir()) / "sgsb_reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)

DEFAULT_OUTLETS = [
    "SBH307",
    "SBH458",
    "SBH001",
    "SBH002",
    "SBH003",
]

# ============================================================
# RUNTIME STATE
# ============================================================

records = db.RecordMap()          # rid -> active/recent record cache
builders = db.SessionMap("builder")         # user_id -> builder state
input_modes = db.SessionMap("input")      # user_id -> remark input state
admin_sessions = db.SessionMap("admin")   # user_id -> admin auth state
report_sessions = db.SessionMap("report")  # user_id -> report query state
media_sessions = db.SessionMap("media")  # durable album association and unassigned uploads
processed_updates = {}

state_lock = threading.RLock()
gist_lock = threading.RLock()
update_locks = tuple(threading.RLock() for _ in range(64))
http_local = threading.local()

OUTLETS = list(DEFAULT_OUTLETS)
OUTLET_STORAGE_OK = False
OUTLET_STORAGE_ERROR = ""

def http_session():
    session = getattr(http_local, "session", None)
    if session is None:
        session = requests.Session()
        http_local.session = session
    return session


def update_lock(update):
    callback = update.get("callback_query") or {}
    message = callback.get("message") or update.get("message") or {}
    actor = callback.get("from") or message.get("from") or {}
    chat = message.get("chat") or {}
    identity = actor.get("id") or chat.get("id") or update.get("update_id") or 0
    return update_locks[int(identity) % len(update_locks)]

# ============================================================
# GENERAL HELPERS
# ============================================================

def now_my():
    return datetime.now(TZ)


def safe_html(value):
    s = str(value or "")
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def normalize_outlet(value):
    value = (value or "").strip().upper()
    value = re.sub(r"[^A-Z0-9_-]", "", value)
    return value[:40]


def normalize_seal(value):
    value = (value or "").strip().upper()
    value = re.sub(r"\s+", " ", value)
    return value[:120]


def hashtag(value):
    return re.sub(
        r"[^A-Z0-9_]",
        "",
        str(value).upper().replace("-", "_").replace(" ", "_")
    )


def encode_seal(seal):
    return base64.urlsafe_b64encode(seal.encode("utf-8")).decode("ascii").rstrip("=")


def decode_seal(value):
    try:
        pad = "=" * ((4 - len(value) % 4) % 4)
        return base64.urlsafe_b64decode((value + pad).encode("ascii")).decode("utf-8")
    except Exception:
        return ""


def make_rid():
    return now_my().strftime("%Y%m%d%H%M%S") + secrets.token_hex(3).upper()


def record_display_id(record):
    return f"{record['outlet']}-{record['rid']}"


def parse_date_input(value):
    value = (value or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    return None


def display_date(value):
    if not value:
        return "-"
    try:
        return date.fromisoformat(str(value)).strftime("%d/%m/%Y")
    except Exception:
        return str(value)


def display_datetime(value):
    if not value:
        return "-"
    try:
        dt = datetime.fromisoformat(str(value))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=TZ)
        else:
            dt = dt.astimezone(TZ)
        return dt.strftime("%d/%m/%Y %I:%M %p")
    except Exception:
        return str(value)


def remember_update(update_id):
    if update_id is None:
        return True

    stamp = time.time()
    with state_lock:
        stale = [k for k, ts in processed_updates.items() if stamp - ts > 3600]
        for key in stale:
            processed_updates.pop(key, None)

        if update_id in processed_updates:
            return False

        processed_updates[update_id] = stamp
        return True


def months_between(start_date, end_date):
    months = []
    current = date(start_date.year, start_date.month, 1)
    final = date(end_date.year, end_date.month, 1)
    while current <= final:
        months.append(current.strftime("%Y%m"))
        if current.month == 12:
            current = date(current.year + 1, 1, 1)
        else:
            current = date(current.year, current.month + 1, 1)
    return months


def cleanup_old_runtime_records(max_age_hours=72):
    pass  # Supabase remains authoritative; never delete permanent records here.

# ============================================================
# TELEGRAM HELPERS
# ============================================================

def telegram(method, payload=None, timeout=30):
    if not BOT_TOKEN:
        return {"ok": False, "description": "BOT_TOKEN missing"}

    try:
        response = http_session().post(
            f"{API}/{method}",
            json=payload or {},
            timeout=timeout
        )
        try:
            return response.json()
        except Exception:
            return {
                "ok": False,
                "description": f"Non-JSON Telegram response HTTP {response.status_code}"
            }
    except Exception as e:
        print("Telegram API network error:", method, type(e).__name__, flush=True)
        return {"ok": False, "ambiguous": True, "description": "Telegram sementara tidak tersedia. Sila cuba semula."}


def send_message(chat_id, text, keyboard=None):
    payload = {
        "chat_id": chat_id,
        "text": text,
        "parse_mode": "HTML"
    }
    if keyboard:
        payload["reply_markup"] = keyboard
    return telegram("sendMessage", payload)


def send_plain_message(chat_id, text, keyboard=None):
    payload = {
        "chat_id": chat_id,
        "text": text
    }
    if keyboard:
        payload["reply_markup"] = keyboard
    return telegram("sendMessage", payload)


def edit_message(chat_id, message_id, text, keyboard=None):
    payload = {
        "chat_id": chat_id,
        "message_id": message_id,
        "text": text,
        "parse_mode": "HTML"
    }
    if keyboard:
        payload["reply_markup"] = keyboard

    result = telegram("editMessageText", payload)
    if (
        not result.get("ok")
        and "message is not modified" in str(result.get("description", "")).lower()
    ):
        return {"ok": True}
    return result


def answer_callback(callback_id, text=None):
    payload = {"callback_query_id": callback_id}
    if text:
        payload["text"] = text[:200]
    return telegram("answerCallbackQuery", payload)


def send_media_group(chat_id, media):
    return telegram(
        "sendMediaGroup",
        {"chat_id": chat_id, "media": media},
        timeout=45
    )


def send_document_file(chat_id, filepath, caption=""):
    try:
        with open(filepath, "rb") as f:
            response = http_session().post(
                f"{API}/sendDocument",
                data={
                    "chat_id": str(chat_id),
                    "caption": caption[:1000]
                },
                files={"document": (Path(filepath).name, f)},
                timeout=90
            )
        return response.json()
    except Exception as e:
        print("sendDocument error:", type(e).__name__, flush=True)
        return {"ok": False, "ambiguous": True, "description": "Telegram sementara tidak tersedia. Sila cuba semula."}


def send_photo_file(chat_id, filepath, caption=""):
    try:
        with open(filepath, "rb") as f:
            response = http_session().post(
                f"{API}/sendPhoto",
                data={
                    "chat_id": str(chat_id),
                    "caption": caption[:1000]
                },
                files={"photo": (Path(filepath).name, f)},
                timeout=90
            )
        return response.json()
    except Exception as e:
        print("sendPhoto error:", type(e).__name__, flush=True)
        return {"ok": False, "ambiguous": True, "description": "Telegram sementara tidak tersedia. Sila cuba semula."}


def send_long_plain(chat_id, text, chunk_size=3500):
    text = str(text or "")
    if not text:
        return

    parts = []
    remaining = text
    while len(remaining) > chunk_size:
        cut = remaining.rfind("\n", 0, chunk_size)
        if cut < chunk_size // 2:
            cut = chunk_size
        parts.append(remaining[:cut])
        remaining = remaining[cut:].lstrip("\n")
    if remaining:
        parts.append(remaining)

    for index, part in enumerate(parts, start=1):
        prefix = f"[{index}/{len(parts)}]\n" if len(parts) > 1 else ""
        send_plain_message(chat_id, prefix + part)

# ============================================================
# SUPABASE STORAGE
# ============================================================

def storage_configured():
    return bool(os.environ.get("SUPABASE_URL") and os.environ.get("SUPABASE_SECRET_KEY"))


def load_outlets_from_supabase():
    global OUTLETS, OUTLET_STORAGE_OK, OUTLET_STORAGE_ERROR
    try:
        data = db.api('sgsb_outlets', params={'active':'eq.true','select':'code','order':'code.asc'})
        OUTLETS = [r['code'] for r in data]
        OUTLET_STORAGE_OK = True
        OUTLET_STORAGE_ERROR = ''
        return True
    except db.StorageError:
        OUTLETS = []
        OUTLET_STORAGE_OK = False
        OUTLET_STORAGE_ERROR = db.FRIENDLY
        return False


def ensure_outlets_available():
    with state_lock:
        ready = OUTLET_STORAGE_OK and bool(OUTLETS)
    return ready or load_outlets_from_supabase()


def save_outlets_to_supabase(outlets):
    db.rpc('sgsb_set_outlets', {'p_codes':sorted(outlets)})
    return True, None


def normalize_record_for_storage(record):
    return dict(record)


def upsert_record_persistent(record):
    db.save(normalize_record_for_storage(record))
    return True, None


def get_record_persistent(rid, bagging_date=None):
    return db.get_record(rid)


def load_records_range(start_date, end_date, outlet=None):
    params={'bagging_date':'gte.'+start_date.isoformat(),'and':'(bagging_date.lte.'+end_date.isoformat()+')'}
    if outlet: params['outlet']='eq.'+outlet
    return db.rows(params)


def load_all_records():
    return db.rows()


def find_existing_record(outlet, seal, bagging_date):
    rows=db.rows({'outlet':'eq.'+outlet,'seal':'eq.'+seal.upper(),'bagging_date':'eq.'+bagging_date,'status':'neq.CANCELLED'})
    return rows[0] if rows else None

# ============================================================
# KEYBOARDS
# ============================================================

def outlet_keyboard(prefix="outlet"):
    with state_lock:
        outlets = sorted(OUTLETS)

    rows = []
    for i in range(0, len(outlets), 2):
        rows.append([
            {"text": outlet, "callback_data": f"{prefix}:{outlet}"}
            for outlet in outlets[i:i + 2]
        ])

    if not rows:
        rows = [[{"text": "Tiada outlet", "callback_data": "noop"}]]
    return {"inline_keyboard": rows}


def record_keyboard(rid, completed=False, failed=False):
    rows = []

    if not completed:
        rows.append([{"text": "📝 Remark", "callback_data": f"remark:{rid}"}])

    if failed:
        rows.append([{"text": "🔄 CUBA ARCHIVE SEMULA", "callback_data": f"retry:{rid}"}])

    rows.append([{"text": "➕ NEW BAGGING", "callback_data": "new"}])

    if not completed:
        rows.append([{"text": "❌ Batal Record", "callback_data": f"cancel_record:{rid}"}])

    return {"inline_keyboard": rows}


def admin_keyboard():
    return {
        "inline_keyboard": [
            [{"text": "📋 Senarai Outlet", "callback_data": "admin:list"}],
            [{"text": "➕ Tambah Outlet", "callback_data": "admin:add"}],
            [{"text": "🗑 Padam Outlet", "callback_data": "admin:delete"}],
            [{"text": "🔄 Sync Outlet Master", "callback_data": "admin:sync"}],
            [{"text": "📊 Report", "callback_data": "report:menu"}],
            [{"text": "🚪 Keluar Admin", "callback_data": "admin:logout"}],
        ]
    }


def delete_outlet_keyboard():
    with state_lock:
        outlets = sorted(OUTLETS)

    rows = [
        [{"text": f"🗑 {outlet}", "callback_data": f"admin:del:{outlet}"}]
        for outlet in outlets
    ]
    rows.append([{"text": "⬅️ Kembali", "callback_data": "admin:menu"}])
    return {"inline_keyboard": rows}


def late_date_keyboard():
    return {
        "inline_keyboard": [
            [
                {"text": "📅 Hari Ini", "callback_data": "late_date:today"},
                {"text": "⏮ Semalam", "callback_data": "late_date:yesterday"},
            ]
        ]
    }


def report_menu_keyboard():
    return {
        "inline_keyboard": [
            [
                {"text": "📅 Hari Ini", "callback_data": "report:q:today"},
                {"text": "⏮ Semalam", "callback_data": "report:q:yesterday"},
            ],
            [{"text": "🗓 Date Range", "callback_data": "report:q:range"}],
            [{"text": "🏢 Outlet + Date Range", "callback_data": "report:q:outlet"}],
            [{"text": "⏳ Semua Incomplete", "callback_data": "report:q:incomplete"}],
            [{"text": "⚠️ Late Upload / Late Complete", "callback_data": "report:q:late"}],
            [{"text": "❌ Tutup", "callback_data": "report:close"}],
        ]
    }


def report_format_keyboard():
    return {
        "inline_keyboard": [
            [
                {"text": "📊 Summary", "callback_data": "report:fmt:summary"},
                {"text": "📋 Full Detail", "callback_data": "report:fmt:detail"},
            ],
            [
                {"text": "📄 PDF", "callback_data": "report:fmt:pdf"},
                {"text": "📗 Excel", "callback_data": "report:fmt:excel"},
            ],
            [
                {"text": "📝 TXT", "callback_data": "report:fmt:txt"},
                {"text": "🖼 Chart", "callback_data": "report:fmt:chart"},
            ],
            [{"text": "📢 Post Summary ke Archive", "callback_data": "report:fmt:archive"}],
            [{"text": "⬅️ Pilih Report Lain", "callback_data": "report:menu"}],
        ]
    }

# ============================================================
# RECORD CARD + RECORD LIFECYCLE
# ============================================================

def record_ref(record):
    return (
        f"SGR5|{record['rid']}|{record['outlet']}|"
        f"{record['bagging_date']}|{encode_seal(record['seal'])}"
    )


def parse_record_ref(text):
    text = text or ""
    match = re.search(
        r"SGR5\|([A-Z0-9]+)\|([A-Z0-9_-]+)\|(\d{4}-\d{2}-\d{2})\|([A-Za-z0-9_-]+)",
        text
    )
    if not match:
        return None

    rid, outlet, bagging_date, seal_encoded = match.groups()
    seal = decode_seal(seal_encoded)
    if not seal:
        return None

    return {
        "rid": rid,
        "outlet": outlet,
        "bagging_date": bagging_date,
        "seal": seal
    }


def record_card_text(record):
    photo_status = "✅" if record.get("photo") else "⏳"
    video_status = "✅" if record.get("video") else "⏳"
    status = record.get("status", "OPEN")
    mode = record.get("entry_mode", "NORMAL")

    if status in ("COMPLETE", "LATE_COMPLETE") and record.get("archived"):
        title = "✅ <b>BAGGING ARCHIVED</b>"
        instruction = "Rekod lengkap dan telah dihantar ke Archive."
    elif status == "ARCHIVE_FAILED" or record.get("archive_failed"):
        title = "⚠️ <b>ARCHIVE GAGAL</b>"
        instruction = "Tekan <b>CUBA ARCHIVE SEMULA</b>."
    elif status == "CANCELLED":
        title = "❌ <b>BAGGING RECORD DIBATALKAN</b>"
        instruction = "Record ini tidak lagi aktif."
    else:
        title = "📦 <b>BAGGING SESSION OPEN</b>"
        instruction = (
            "↩️ <b>REPLY Record Card ini</b> dengan 1 gambar seal dan 1 video bagging.\n\n"
            "Anda boleh tekan <b>NEW BAGGING</b> sekarang dan buka Seal seterusnya "
            "walaupun video record ini masih uploading."
        )

    late_line = ""
    if mode == "LATE":
        late_line = (
            f"\n⚠️ Mode: <b>LATE / BACKDATED</b>\n"
            f"📝 Late Reason: {safe_html(record.get('late_reason') or '—')}"
        )

    return (
        f"{title}\n\n"
        f"🆔 <code>{safe_html(record_display_id(record))}</code>\n"
        f"🏢 Outlet: <b>{safe_html(record['outlet'])}</b>\n"
        f"🔒 Seal: <b>{safe_html(record['seal'])}</b>\n"
        f"📅 Bagging Date: <b>{display_date(record.get('bagging_date'))}</b>"
        f"{late_line}\n\n"
        f"📷 Gambar: {photo_status}\n"
        f"🎥 Video: {video_status}\n"
        f"📝 Remark: {safe_html(record.get('remark') or '—')}\n\n"
        f"{instruction}\n\n"
        f"<code>{safe_html(record_ref(record))}</code>"
    )


def update_record_card(record):
    message_id = record.get("card_message_id")
    chat_id = record.get("chat_id")
    if not message_id or not chat_id:
        return

    completed = record.get("status") in ("COMPLETE", "LATE_COMPLETE", "CANCELLED")
    failed = record.get("status") == "ARCHIVE_FAILED" or record.get("archive_failed")

    edit_message(
        chat_id,
        message_id,
        record_card_text(record),
        record_keyboard(record["rid"], completed=completed, failed=failed)
    )


def send_new_record_card(record):
    result = send_message(
        record["chat_id"],
        record_card_text(record),
        record_keyboard(record["rid"])
    )
    if result.get("ok"):
        record["card_message_id"] = result.get("result", {}).get("message_id")
    return result


def create_record(user_id, chat_id, outlet, seal, bagging_date, entry_mode="NORMAL", late_reason=""):
    existing = find_existing_record(outlet, seal, bagging_date)
    if existing:
        return None, {"duplicate": True, "record": existing}

    record = {
        "rid": make_rid(),
        "owner_id": user_id,
        "chat_id": chat_id,
        "outlet": outlet,
        "seal": seal,
        "bagging_date": bagging_date,
        "entry_mode": entry_mode,
        "late_reason": late_reason[:500],
        "photo": None,
        "video": None,
        "remark": "",
        "card_message_id": None,
        "created_at": now_my().isoformat(),
        "completed_at": None,
        "archived_at": None,
        "archived": False,
        "archive_failed": False,
        "archiving": False,
        "status": "OPEN",
        "cancelled_at": None
    }

    try:
        db.api('sgsb_records','POST',{'rid':record['rid'],'data':record})
    except db.DuplicateRecord:
        return None, {'duplicate':True,'record':find_existing_record(outlet,seal,bagging_date)}

    result = send_new_record_card(record)
    if not result.get("ok"):
        return None, result

    ok, error = upsert_record_persistent(record)
    if not ok:
        edit_message(
            chat_id,
            record.get("card_message_id"),
            "❌ <b>STORAGE ERROR</b>\n\n"
            "Bagging Session tidak dibuka kerana metadata tidak dapat disimpan secara permanent.\n"
            f"<code>{safe_html(error)}</code>"
        )
        return None, {"ok": False, "description": error}

    with state_lock:
        records[record["rid"]] = record

    return record, {"ok": True}


def resume_existing_record(user_id, chat_id, stored):
    record = dict(stored)
    if record.get("owner_id") != user_id and not admin_authenticated(user_id):
        return None, "Rekod ini milik pengguna lain. Hubungi admin untuk recovery."
    record["owner_id"] = user_id
    record["chat_id"] = chat_id
    record["archiving"] = False

    if record.get("status") in ("COMPLETE", "LATE_COMPLETE"):
        return None, "Rekod ini sudah COMPLETE dan berada dalam Archive."
    if record.get("status") == "CANCELLED":
        return None, "Rekod ini sudah dibatalkan."

    result = send_new_record_card(record)
    if not result.get("ok"):
        return None, result.get("description", "Gagal membuka semula Record Card")

    ok, error = upsert_record_persistent(record)
    if not ok:
        return None, error

    with state_lock:
        records[record["rid"]] = record

    return record, None


def reconstruct_record_from_reply(user_id, chat_id, reply_message):
    ref = parse_record_ref(reply_message.get("text") or reply_message.get("caption") or "")
    if not ref:
        return None

    rid = ref["rid"]

    with state_lock:
        existing = records.get(rid)
        if existing:
            return existing if existing.get("owner_id") == user_id or admin_authenticated(user_id) else None

    stored = get_record_persistent(rid, ref.get("bagging_date"))
    if not stored or (stored.get("owner_id") != user_id and not admin_authenticated(user_id)):
        return None

    record = dict(stored)
    record["owner_id"] = user_id
    record["chat_id"] = chat_id
    record["card_message_id"] = reply_message.get("message_id")
    record["archiving"] = False

    with state_lock:
        records[rid] = record

    return record


def open_records_for_user(user_id):
    return db.rows({'owner_id':'eq.'+str(user_id),'status':'not.in.(COMPLETE,LATE_COMPLETE,CANCELLED)'})


def resolve_record_for_media(message, user_id, chat_id):
    reply = message.get("reply_to_message")

    if reply:
        record = reconstruct_record_from_reply(user_id, chat_id, reply)
        if record:
            record["owner_id"] = user_id
            record["chat_id"] = chat_id
            return record, None
        return None, "⚠️ Record Card tidak sah atau milik pengguna lain. Reply Record Card yang betul."

    open_records = open_records_for_user(user_id)

    if len(open_records) == 1:
        return open_records[0], None

    if len(open_records) > 1:
        return None, (
            "⚠️ Anda mempunyai lebih daripada satu Bagging Session yang masih OPEN.\n\n"
            "Untuk elakkan video tersalah Seal, <b>REPLY Record Card</b> yang betul kemudian hantar media."
        )

    return None, (
        "⚠️ Tiada Bagging Session yang aktif.\n"
        "Tekan /start atau /late."
    )


def media_session_key(message, user_id, chat_id):
    group = message.get("media_group_id")
    source = "album:" + str(group) if group else "message:" + str(message["message_id"])
    return hashlib.sha256(f"{user_id}:{chat_id}:{source}".encode()).hexdigest()[:16]


def media_reply(record, message):
    return {**message, "reply_to_message": {
        "message_id": record.get("card_message_id"), "text": record_ref(record)
    }}


def choose_pending_media(callback, token, rid, user_id, chat_id):
    pending = media_sessions.get(token)
    if not pending or pending.get("owner_id") != user_id or pending.get("chat_id") != chat_id:
        send_message(chat_id, "⚠️ Pilihan media tidak sah atau milik pengguna lain.")
        return
    if pending.get("rid") and pending["rid"] != rid:
        send_message(chat_id, "⚠️ Album ini sudah dipadankan kepada Seal lain.")
        return
    record = get_record_persistent(rid)
    if not record or record.get("owner_id") != user_id or record.get("chat_id") != chat_id:
        send_message(chat_id, "⚠️ Record Card tidak sah atau milik pengguna lain.")
        return
    if record.get("status") == "CANCELLED" or (
        record.get("status") in ("COMPLETE", "LATE_COMPLETE") and not pending.get("rid")
    ):
        send_message(chat_id, "⚠️ Seal ini sudah selesai atau dibatalkan. Pilih Record Card yang masih OPEN.")
        return
    pending["rid"] = rid  # persist before processing; later album items use the same Seal
    for item in list(pending.get("messages", [])):
        handle_message(media_reply(record, item))
    pending["messages"] = []
    if pending.get("prompt_id"):
        edit_message(chat_id, pending["prompt_id"],
                     "✅ Media dipadankan kepada Seal <b>" + safe_html(record["seal"]) + "</b>.")


def prepare_media(message, user_id, chat_id):
    token = media_session_key(message, user_id, chat_id)
    pending = media_sessions.get(token)
    record, error = resolve_record_for_media(message, user_id, chat_id)
    if pending and pending.get("rid"):
        bound = get_record_persistent(pending["rid"])
        if message.get("reply_to_message") and (not record or record["rid"] != pending["rid"]):
            send_message(chat_id, "⚠️ Album ini sudah dipadankan kepada Seal lain. Hantar media berasingan untuk Seal seterusnya.")
            return None
        if not bound:
            send_message(chat_id, "⚠️ Record album tidak dijumpai. Hubungi admin.")
            return None
        return bound
    if pending and not message.get("reply_to_message"):
        record = None  # An initially ambiguous album needs an explicit choice, even if another session closes.
    if record:
        if message.get("media_group_id") or pending:
            if not pending:
                media_sessions[token] = {"owner_id": user_id, "chat_id": chat_id,
                                         "messages": [], "created_at": now_my().isoformat()}
                pending = media_sessions[token]
            pending["rid"] = record["rid"]
            for item in list(pending.get("messages", [])):
                handle_message(media_reply(record, item))
            pending["messages"] = []
            record = get_record_persistent(record["rid"])
        return record
    candidates = open_records_for_user(user_id) if not message.get("reply_to_message") else []
    if (len(candidates) <= 1 and not pending) or not ("photo" in message or "video" in message):
        send_message(chat_id, error)
        return None
    if not pending:
        media_sessions[token] = {"owner_id": user_id, "chat_id": chat_id,
                                 "messages": [], "created_at": now_my().isoformat()}
        pending = media_sessions[token]
    items = list(pending.get("messages", []))
    if not any(item["message_id"] == message["message_id"] for item in items):
        items.append(message)
        pending["messages"] = items  # file_id survives restarts; no re-upload needed
    if not pending.get("prompt_id"):
        keyboard = {"inline_keyboard": [[{
            "text": (r["outlet"] + " · " + r["seal"] + " · " + r["bagging_date"])[:100],
            "callback_data": "media:" + token + ":" + r["rid"]
        }] for r in candidates[:90]]}
        result = send_message(chat_id,
            "📎 Media diterima tetapi belum dipadankan kerana ada beberapa sesi OPEN.\n\n"
            "<b>Pilih Seal yang betul untuk media/album ini.</b> Tidak perlu hantar semula.\n"
            "Untuk upload seterusnya, reply Record Card yang betul.", keyboard)
        if result.get("ok"):
            pending["prompt_id"] = result["result"]["message_id"]
    return None


def determine_final_status(record):
    try:
        bag_date = date.fromisoformat(record["bagging_date"])
    except Exception:
        bag_date = now_my().date()

    completed = record.get("completed_at")
    try:
        completed_date = datetime.fromisoformat(completed).astimezone(TZ).date()
    except Exception:
        completed_date = now_my().date()

    is_late = (
        record.get("entry_mode") == "LATE"
        or completed_date > bag_date
    )
    return "LATE_COMPLETE" if is_late else "COMPLETE"

# ============================================================
# ARCHIVE — ONE SEAL = ONE TELEGRAM ALBUM
# ============================================================

def archive_caption(record):
    is_late = determine_final_status(record) == "LATE_COMPLETE"
    heading = "⚠️ <b>LATE BAGGING UPLOAD</b>" if is_late else "📦 <b>SGSB BAGGING RECORD</b>"

    late_text = ""
    if is_late:
        late_text = (
            f"📅 Bagging Date: <b>{display_date(record.get('bagging_date'))}</b>\n"
            f"📤 Completed: <b>{display_datetime(record.get('completed_at'))}</b>\n"
            f"📝 Late Reason: {safe_html(record.get('late_reason') or '—')}\n"
        )

    return (
        f"{heading}\n"
        "━━━━━━━━━━━━━━━━━━\n"
        f"🆔 <code>{safe_html(record_display_id(record))}</code>\n"
        f"🏢 <b>{safe_html(record['outlet'])}</b>\n"
        f"🔒 <b>{safe_html(record['seal'])}</b>\n"
        f"{late_text}"
        f"📝 Remark: {safe_html(record.get('remark') or '—')}\n"
        "━━━━━━━━━━━━━━━━━━\n"
        "📷 1 Photo • 🎥 1 Video\n\n"
        f"#{hashtag(record['outlet'])} "
        f"#{str(record.get('bagging_date', '')).replace('-', '')} "
        f"#SEAL_{hashtag(record['seal'])}"
    )


def archive_record(record):
    if record.get("archived"):
        return True, None

    if not record.get("photo") or not record.get("video"):
        return False, "Record belum lengkap"

    if record.get("archive_delivery_state") in ("SENDING", "UNKNOWN"):
        return False, "Status Archive perlu disemak oleh admin sebelum retry bagi mengelakkan album duplicate."
    if record.get("archiving"):
        return False, "Archive sedang diproses"

    record["archiving"] = True
    record["archive_delivery_state"] = "SENDING"
    upsert_record_persistent(record)
    record["archive_failed"] = False
    update_record_card(record)

    media = [
        {
            "type": "photo",
            "media": record["photo"]["file_id"],
            "caption": archive_caption(record),
            "parse_mode": "HTML"
        },
        {
            "type": "video",
            "media": record["video"]["file_id"]
        }
    ]

    result = send_media_group(ARCHIVE_CHAT_ID, media)
    record["archiving"] = False

    if not result.get("ok"):
        record["archive_delivery_state"] = "UNKNOWN" if result.get("ambiguous") else "FAILED"
        record["archive_failed"] = True
        record["status"] = "ARCHIVE_FAILED"
        upsert_record_persistent(record)
        update_record_card(record)
        error = result.get("description", str(result))
        print("Archive failed:", record_display_id(record), error, flush=True)
        return False, error

    record["archived"] = True
    record["archive_delivery_state"] = "SENT"
    record["archive_message_ids"] = [m["message_id"] for m in result.get("result", [])]
    record["archive_failed"] = False
    record["archived_at"] = now_my().isoformat()
    record["status"] = determine_final_status(record)

    ok, storage_error = upsert_record_persistent(record)
    update_record_card(record)

    if not ok:
        send_message(
            record["chat_id"],
            "⚠️ <b>ARCHIVE BERJAYA, TETAPI REPORT LOG GAGAL DISIMPAN</b>\n\n"
            f"Record: <code>{safe_html(record_display_id(record))}</code>\n"
            f"Error: <code>{safe_html(storage_error)}</code>"
        )
    else:
        send_message(
            record["chat_id"],
            "✅ <b>ARCHIVE BERJAYA</b>\n\n"
            f"🆔 <code>{safe_html(record_display_id(record))}</code>\n"
            f"🏢 <b>{safe_html(record['outlet'])}</b>\n"
            f"🔒 <b>{safe_html(record['seal'])}</b>\n"
            f"📊 Status: <b>{record['status']}</b>\n\n"
            "Gambar + video telah dihantar sebagai satu album.",
            {"inline_keyboard": [[{"text": "➕ NEW BAGGING", "callback_data": "new"}]]}
        )

    return True, None


def maybe_archive_record(record):
    if record.get("photo") and record.get("video") and not record.get("archived"):
        if not record.get("completed_at"):
            record["completed_at"] = now_my().isoformat()
        record["status"] = "READY"

        ok, error = upsert_record_persistent(record)
        if not ok:
            send_message(
                record["chat_id"],
                "⚠️ <b>MEDIA LENGKAP, TETAPI METADATA GAGAL DISIMPAN</b>\n"
                "Archive ditangguhkan supaya report tidak kehilangan rekod.\n"
                f"<code>{safe_html(error)}</code>"
            )
            return False, error

        return archive_record(record)

    upsert_record_persistent(record)
    update_record_card(record)
    return True, None

# ============================================================
# STAFF BUILDERS: NORMAL + LATE
# ============================================================

def start_builder(chat_id, user_id, mode="NORMAL"):
    if not load_outlets_from_supabase():
        raise db.StorageError(db.FRIENDLY)
    with state_lock:
        builders[user_id] = {
            "mode": mode,
            "step": "outlet",
            "outlet": None,
            "seal": None,
            "bagging_date": None,
            "late_reason": ""
        }

    title = "⚠️ <b>LATE / BACKDATED BAGGING</b>" if mode == "LATE" else "📦 <b>NEW BAGGING</b>"
    send_message(
        chat_id,
        f"{title}\n\nPilih <b>Outlet Code</b>:",
        outlet_keyboard("late_outlet" if mode == "LATE" else "outlet")
    )


def complete_normal_builder(chat_id, user_id, builder, seal):
    record, result = create_record(
        user_id,
        chat_id,
        builder["outlet"],
        seal,
        now_my().date().isoformat(),
        entry_mode="NORMAL"
    )

    if not record:
        if result.get("duplicate"):
            existing = result["record"]
            send_message(
                chat_id,
                "⚠️ <b>DUPLICATE RECORD DIKESAN</b>\n\n"
                f"Outlet: <b>{safe_html(existing.get('outlet'))}</b>\n"
                f"Seal: <b>{safe_html(existing.get('seal'))}</b>\n"
                f"Date: <b>{display_date(existing.get('bagging_date'))}</b>\n"
                f"Status: <b>{safe_html(existing.get('status'))}</b>\n\n"
                "Gunakan /late jika mahu sambung record incomplete lama."
            )
        else:
            send_message(
                chat_id,
                "❌ Gagal membuka Bagging Session.\n"
                f"<code>{safe_html(result.get('description', result))}</code>"
            )
        return

    with state_lock:
        builders.pop(user_id, None)

    send_message(
        chat_id,
        "✅ <b>SESSION DIBUKA</b>\n\n"
        "Reply Record Card di atas dengan gambar/video.\n\n"
        "Anda boleh tekan <b>NEW BAGGING</b> terus untuk Seal seterusnya sementara video sebelumnya masih uploading."
    )


def handle_late_date(chat_id, user_id, selected_date):
    builder = builders.get(user_id)
    if not builder or builder.get("mode") != "LATE":
        send_message(chat_id, "⚠️ Late session sudah tamat. Taip /late semula.")
        return

    builder["bagging_date"] = selected_date.isoformat()

    existing = find_existing_record(
        builder["outlet"],
        builder["seal"],
        builder["bagging_date"]
    )

    if existing:
        record, error = resume_existing_record(user_id, chat_id, existing)
        builders.pop(user_id, None)

        if record:
            missing = []
            if not record.get("photo"):
                missing.append("gambar")
            if not record.get("video"):
                missing.append("video")
            send_message(
                chat_id,
                "✅ <b>REKOD LAMA DIJUMPAI</b>\n\n"
                f"🆔 <code>{safe_html(record_display_id(record))}</code>\n"
                f"📅 Bagging Date: <b>{display_date(record.get('bagging_date'))}</b>\n"
                f"⏳ Masih perlu: <b>{safe_html(', '.join(missing) or 'Archive retry')}</b>\n\n"
                "Reply Record Card yang baru dihantar untuk lengkapkan media yang tertinggal."
            )
        else:
            send_message(chat_id, f"ℹ️ {safe_html(error)}")
        return

    builder["step"] = "late_reason"
    send_message(
        chat_id,
        "📝 <b>Sebab Late Upload</b>\n\n"
        "Taip sebab, contoh: <i>Staff terlupa upload video</i>.\n"
        "Jika tiada sebab, taip <code>/skip</code>."
    )


def finalize_late_builder(chat_id, user_id, reason):
    builder = builders.get(user_id)
    if not builder:
        return

    record, result = create_record(
        user_id,
        chat_id,
        builder["outlet"],
        builder["seal"],
        builder["bagging_date"],
        entry_mode="LATE",
        late_reason=reason
    )

    builders.pop(user_id, None)

    if not record:
        send_message(
            chat_id,
            "❌ Gagal membuka Late Bagging Session.\n"
            f"<code>{safe_html(result.get('description', result))}</code>"
        )
        return

    send_message(
        chat_id,
        "✅ <b>LATE SESSION DIBUKA</b>\n\n"
        f"📅 Bagging Date: <b>{display_date(record['bagging_date'])}</b>\n"
        "Reply Record Card dengan media yang tertinggal."
    )

# ============================================================
# ADMIN
# ============================================================

def begin_admin(chat_id, user_id):
    admin_sessions[user_id] = {"authenticated": False, "step": "passcode"}
    send_message(chat_id, "🔐 <b>ADMIN OUTLET MANAGEMENT</b>\n\nMasukkan Admin Passcode.")


def admin_authenticated(user_id):
    return bool(admin_sessions.get(user_id, {}).get("authenticated"))


def show_admin_menu(chat_id, user_id):
    admin_sessions[user_id] = {"authenticated": True, "step": "menu"}
    storage_text = "✅ Permanent Storage: Connected" if OUTLET_STORAGE_OK else "⚠️ Permanent Storage: NOT CONNECTED"
    send_message(
        chat_id,
        "⚙️ <b>ADMIN OUTLET MANAGEMENT</b>\n\n"
        f"{storage_text}\n🏢 Total Outlet: {len(OUTLETS)}",
        admin_keyboard()
    )


def handle_admin_text(chat_id, user_id, text):
    adm = admin_sessions.get(user_id)
    if not adm:
        return False

    if adm.get("step") == "passcode":
        if ADMIN_PASSCODE and secrets.compare_digest(text, ADMIN_PASSCODE):
            show_admin_menu(chat_id, user_id)
        else:
            admin_sessions.pop(user_id, None)
            send_message(chat_id, "❌ Passcode salah.")
        return True

    if not adm.get("authenticated"):
        return False

    if adm.get("step") == "add":
        outlet = normalize_outlet(text)
        if not outlet:
            send_message(chat_id, "⚠️ Outlet Code tidak sah.")
            return True

        with state_lock:
            current = list(OUTLETS)

        if outlet in current:
            adm["step"] = "menu"
            send_message(chat_id, f"ℹ️ <b>{safe_html(outlet)}</b> sudah wujud.", admin_keyboard())
            return True

        updated = current + [outlet]
        ok, error = save_outlets_to_supabase(updated)
        if not ok:
            adm["step"] = "menu"
            send_message(
                chat_id,
                "❌ <b>OUTLET TIDAK DISIMPAN</b>\n"
                f"<code>{safe_html(error)}</code>",
                admin_keyboard()
            )
            return True

        with state_lock:
            OUTLETS[:] = sorted(updated)

        adm["step"] = "menu"
        send_message(
            chat_id,
            f"✅ Outlet <b>{safe_html(outlet)}</b> disimpan secara permanent.",
            admin_keyboard()
        )
        return True

    return False

# ============================================================
# REPORT ENGINE
# ============================================================

def begin_report(chat_id, user_id):
    if not admin_authenticated(user_id):
        report_sessions[user_id] = {"step": "passcode"}
        send_message(
            chat_id,
            "🔐 <b>SGSB REPORT</b>\n\nMasukkan Admin Passcode."
        )
        return

    report_sessions[user_id] = {"step": "menu"}
    send_message(
        chat_id,
        "📊 <b>SGSB BAGGING REPORT</b>\n\nPilih jenis report:",
        report_menu_keyboard()
    )


def report_status_label(record):
    status = record.get("status") or "OPEN"
    if status in ("OPEN", "READY"):
        if not record.get("photo") and not record.get("video"):
            return "INCOMPLETE (NO MEDIA)"
        if not record.get("photo"):
            return "INCOMPLETE (MISSING PHOTO)"
        if not record.get("video"):
            return "INCOMPLETE (MISSING VIDEO)"
        return "INCOMPLETE"
    return status


def build_report_context(rows, start_date=None, end_date=None, outlet=None, title="SGSB BAGGING REPORT"):
    active = [r for r in rows if r.get("status") != "CANCELLED"]
    cancelled = [r for r in rows if r.get("status") == "CANCELLED"]

    on_time = [r for r in active if r.get("status") == "COMPLETE"]
    late_complete = [r for r in active if r.get("status") == "LATE_COMPLETE"]
    archive_failed = [r for r in active if r.get("status") == "ARCHIVE_FAILED"]
    incomplete = [
        r for r in active
        if r.get("status") in (None, "OPEN", "READY")
        or (not r.get("photo") or not r.get("video"))
    ]

    outlet_counts = Counter(r.get("outlet", "UNKNOWN") for r in active)
    daily_counts = Counter(r.get("bagging_date", "UNKNOWN") for r in active)

    submitted_outlets = {r.get("outlet") for r in active if r.get("outlet")}
    with state_lock:
        master = set(OUTLETS)
    no_submission = sorted(master - submitted_outlets) if not outlet else []

    context = {
        "title": title,
        "generated_at": now_my().isoformat(),
        "start_date": start_date.isoformat() if isinstance(start_date, date) else start_date,
        "end_date": end_date.isoformat() if isinstance(end_date, date) else end_date,
        "outlet_filter": outlet,
        "rows": rows,
        "active": active,
        "cancelled": cancelled,
        "on_time": on_time,
        "late_complete": late_complete,
        "archive_failed": archive_failed,
        "incomplete": incomplete,
        "outlet_counts": outlet_counts,
        "daily_counts": daily_counts,
        "submitted_outlets": sorted(submitted_outlets),
        "no_submission": no_submission,
        "missing_photo": sum(1 for r in active if not r.get("photo")),
        "missing_video": sum(1 for r in active if not r.get("video")),
    }
    return context


def report_period_text(context):
    start = context.get("start_date")
    end = context.get("end_date")
    if start and end:
        return f"{display_date(start)} - {display_date(end)}"
    return "All available records"


def report_summary_text(context, html_mode=True):
    esc = safe_html if html_mode else str
    title = context["title"]
    period = report_period_text(context)
    outlet_filter = context.get("outlet_filter") or "Semua Outlet"

    lines = [
        f"📊 {'<b>' if html_mode else ''}{esc(title)}{'</b>' if html_mode else ''}",
        "",
        f"📅 Period: {esc(period)}",
        f"🏢 Filter: {esc(outlet_filter)}",
        f"🕒 Generated: {esc(display_datetime(context['generated_at']))}",
        "",
        "📌 OVERALL",
        f"Total Records: {len(context['active'])}",
        f"✅ Complete On Time: {len(context['on_time'])}",
        f"⚠️ Late Complete: {len(context['late_complete'])}",
        f"⏳ Incomplete: {len(context['incomplete'])}",
        f"❌ Archive Failed: {len(context['archive_failed'])}",
        f"🚫 Cancelled: {len(context['cancelled'])}",
        f"📷 Missing Photo: {context['missing_photo']}",
        f"🎥 Missing Video: {context['missing_video']}",
        f"🏢 Outlet Submitted: {len(context['submitted_outlets'])}",
    ]

    lines += ["", "🏢 OUTLET BREAKDOWN"]
    for outlet, count in sorted(context["outlet_counts"].items(), key=lambda x: (-x[1], x[0])):
        lines.append(f"{esc(outlet)} — {count}")

    lines += ["", "📆 DAILY BREAKDOWN"]
    for day, count in sorted(context["daily_counts"].items()):
        lines.append(f"{esc(display_date(day))} — {count}")

    if context["no_submission"]:
        lines += ["", "⚠️ OUTLET TIADA SUBMISSION"]
        for outlet in context["no_submission"]:
            lines.append(f"• {esc(outlet)}")

    if context["incomplete"]:
        lines += ["", "⏳ INCOMPLETE RECORDS"]
        for r in context["incomplete"][:50]:
            lines.append(
                f"• {esc(r.get('outlet'))} | {esc(r.get('seal'))} | "
                f"{esc(display_date(r.get('bagging_date')))} | {esc(report_status_label(r))}"
            )
        if len(context["incomplete"]) > 50:
            lines.append(f"... +{len(context['incomplete']) - 50} lagi")

    if context["late_complete"]:
        lines += ["", "⚠️ LATE COMPLETE"]
        for r in context["late_complete"][:50]:
            lines.append(
                f"• {esc(r.get('outlet'))} | {esc(r.get('seal'))} | "
                f"Bagging {esc(display_date(r.get('bagging_date')))} | "
                f"Completed {esc(display_datetime(r.get('completed_at')))}"
            )
        if len(context["late_complete"]) > 50:
            lines.append(f"... +{len(context['late_complete']) - 50} lagi")

    return "\n".join(lines)


def report_detail_text(context):
    lines = [
        context["title"],
        f"Period: {report_period_text(context)}",
        f"Filter: {context.get('outlet_filter') or 'Semua Outlet'}",
        f"Generated: {display_datetime(context['generated_at'])}",
        "=" * 72,
        "",
        report_summary_text(context, html_mode=False),
        "",
        "=" * 72,
        "FULL DETAIL",
        "=" * 72,
    ]

    for index, r in enumerate(context["rows"], start=1):
        lines += [
            f"{index}. {r.get('outlet', '-')} | {r.get('seal', '-')}",
            f"   Record ID   : {r.get('outlet', '-')}-{r.get('rid', '-')}",
            f"   Bagging Date: {display_date(r.get('bagging_date'))}",
            f"   Created     : {display_datetime(r.get('created_at'))}",
            f"   Completed   : {display_datetime(r.get('completed_at'))}",
            f"   Archived    : {display_datetime(r.get('archived_at'))}",
            f"   Status      : {report_status_label(r)}",
            f"   Photo       : {'YES' if r.get('photo') else 'NO'}",
            f"   Video       : {'YES' if r.get('video') else 'NO'}",
            f"   Remark      : {r.get('remark') or '-'}",
            f"   Late Reason : {r.get('late_reason') or '-'}",
            "",
        ]

    return "\n".join(lines)


def safe_filename_part(value):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(value or "report")).strip("_")[:80]


def report_base_filename(context):
    period = report_period_text(context).replace("/", "-").replace(" ", "_")
    outlet = context.get("outlet_filter") or "ALL"
    stamp = now_my().strftime("%Y%m%d_%H%M%S")
    return f"SGSB_BAGGING_{safe_filename_part(outlet)}_{safe_filename_part(period)}_{stamp}"


def generate_txt_report(context):
    path = REPORT_DIR / f"{report_base_filename(context)}.txt"
    path.write_text(report_detail_text(context), encoding="utf-8")
    return str(path)


def generate_excel_report(context):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    path = REPORT_DIR / f"{report_base_filename(context)}.xlsx"
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    summary_rows = [
        ["SGSB BAGGING REPORT", ""],
        ["Period", report_period_text(context)],
        ["Outlet Filter", context.get("outlet_filter") or "Semua Outlet"],
        ["Generated", display_datetime(context["generated_at"])],
        ["Total Records", len(context["active"])],
        ["Complete On Time", len(context["on_time"])],
        ["Late Complete", len(context["late_complete"])],
        ["Incomplete", len(context["incomplete"])],
        ["Archive Failed", len(context["archive_failed"])],
        ["Cancelled", len(context["cancelled"])],
        ["Missing Photo", context["missing_photo"]],
        ["Missing Video", context["missing_video"]],
        ["Outlet Submitted", len(context["submitted_outlets"])],
    ]
    for row in summary_rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=16)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 42

    ws_outlet = wb.create_sheet("Outlet Breakdown")
    ws_outlet.append(["Outlet", "Total Records"])
    for outlet, count in sorted(context["outlet_counts"].items(), key=lambda x: (-x[1], x[0])):
        ws_outlet.append([outlet, count])
    ws_outlet.freeze_panes = "A2"
    ws_outlet.auto_filter.ref = ws_outlet.dimensions
    ws_outlet.column_dimensions["A"].width = 24
    ws_outlet.column_dimensions["B"].width = 18

    ws_daily = wb.create_sheet("Daily Breakdown")
    ws_daily.append(["Date", "Total Records"])
    for day, count in sorted(context["daily_counts"].items()):
        ws_daily.append([display_date(day), count])
    ws_daily.freeze_panes = "A2"
    ws_daily.auto_filter.ref = ws_daily.dimensions
    ws_daily.column_dimensions["A"].width = 18
    ws_daily.column_dimensions["B"].width = 18

    ws_records = wb.create_sheet("Records")
    headers = [
        "Record ID", "Outlet", "Seal", "Bagging Date", "Created At",
        "Completed At", "Archived At", "Status", "Photo", "Video",
        "Remark", "Late Reason", "Entry Mode"
    ]
    ws_records.append(headers)

    for r in context["rows"]:
        ws_records.append([
            f"{r.get('outlet', '-')}-{r.get('rid', '-')}",
            r.get("outlet", "-"),
            r.get("seal", "-"),
            display_date(r.get("bagging_date")),
            display_datetime(r.get("created_at")),
            display_datetime(r.get("completed_at")),
            display_datetime(r.get("archived_at")),
            report_status_label(r),
            "YES" if r.get("photo") else "NO",
            "YES" if r.get("video") else "NO",
            r.get("remark") or "-",
            r.get("late_reason") or "-",
            r.get("entry_mode") or "NORMAL",
        ])

    for cell in ws_records[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    ws_records.freeze_panes = "A2"
    ws_records.auto_filter.ref = ws_records.dimensions

    widths = [28, 20, 24, 16, 22, 22, 22, 24, 10, 10, 40, 40, 16]
    for i, width in enumerate(widths, start=1):
        ws_records.column_dimensions[get_column_letter(i)].width = width

    wb.save(path)
    return str(path)


def generate_pdf_report(context):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

    path = REPORT_DIR / f"{report_base_filename(context)}.pdf"
    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCenter",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=18,
        spaceAfter=12
    )

    story = [
        Paragraph(html.escape(context["title"]), title_style),
        Paragraph(
            f"Period: {html.escape(report_period_text(context))} &nbsp;&nbsp; | &nbsp;&nbsp; "
            f"Outlet: {html.escape(context.get('outlet_filter') or 'Semua Outlet')} &nbsp;&nbsp; | &nbsp;&nbsp; "
            f"Generated: {html.escape(display_datetime(context['generated_at']))}",
            styles["Normal"]
        ),
        Spacer(1, 12),
    ]

    summary_data = [
        ["Metric", "Value"],
        ["Total Records", len(context["active"])],
        ["Complete On Time", len(context["on_time"])],
        ["Late Complete", len(context["late_complete"])],
        ["Incomplete", len(context["incomplete"])],
        ["Archive Failed", len(context["archive_failed"])],
        ["Cancelled", len(context["cancelled"])],
        ["Missing Photo", context["missing_photo"]],
        ["Missing Video", context["missing_video"]],
    ]
    t = Table(summary_data, colWidths=[180, 90])
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
    ]))
    story += [t, Spacer(1, 18)]

    outlet_data = [["Outlet", "Records"]] + [
        [outlet, count]
        for outlet, count in sorted(context["outlet_counts"].items(), key=lambda x: (-x[1], x[0]))
    ]
    if len(outlet_data) > 1:
        story.append(Paragraph("Outlet Breakdown", styles["Heading2"]))
        t2 = Table(outlet_data, repeatRows=1, colWidths=[180, 90])
        t2.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ]))
        story += [t2, Spacer(1, 16)]

    story.append(PageBreak())
    story.append(Paragraph("Detailed Records", styles["Heading2"]))

    detail_data = [[
        "#", "Date", "Outlet", "Seal", "Status", "Photo", "Video",
        "Completed", "Remark / Late Reason"
    ]]

    for i, r in enumerate(context["rows"], start=1):
        note = r.get("remark") or "-"
        if r.get("late_reason"):
            note += f" | Late: {r.get('late_reason')}"
        detail_data.append([
            i,
            display_date(r.get("bagging_date")),
            r.get("outlet", "-"),
            r.get("seal", "-"),
            report_status_label(r),
            "YES" if r.get("photo") else "NO",
            "YES" if r.get("video") else "NO",
            display_datetime(r.get("completed_at")),
            Paragraph(html.escape(str(note)), styles["BodyText"]),
        ])

    detail_table = Table(
        detail_data,
        repeatRows=1,
        colWidths=[24, 62, 72, 85, 105, 38, 38, 105, 170]
    )
    detail_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(detail_table)

    doc.build(story)
    return str(path)


def generate_chart_report(context):
    from PIL import Image, ImageDraw, ImageFont

    path = REPORT_DIR / f"{report_base_filename(context)}.png"
    width, height = 1400, 1000
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    def font(size, bold=False):
        candidates = [
            "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        ]
        for candidate in candidates:
            try:
                return ImageFont.truetype(candidate, size)
            except Exception:
                pass
        return ImageFont.load_default()

    title_font = font(34, True)
    heading_font = font(22, True)
    normal_font = font(18, False)
    small_font = font(15, False)

    draw.text((60, 40), context["title"], fill="black", font=title_font)
    draw.text((60, 90), f"Period: {report_period_text(context)}", fill="black", font=normal_font)
    draw.text((60, 120), f"Outlet: {context.get('outlet_filter') or 'Semua Outlet'}", fill="black", font=normal_font)

    summary = [
        ("Total", len(context["active"])),
        ("On Time", len(context["on_time"])),
        ("Late", len(context["late_complete"])),
        ("Incomplete", len(context["incomplete"])),
        ("Archive Failed", len(context["archive_failed"])),
    ]

    x = 60
    y = 175
    for label, value in summary:
        draw.rounded_rectangle((x, y, x + 220, y + 95), radius=12, outline="black", width=2)
        draw.text((x + 16, y + 14), label, fill="black", font=small_font)
        draw.text((x + 16, y + 45), str(value), fill="black", font=heading_font)
        x += 250

    draw.text((60, 315), "Top Outlet by Bagging Records", fill="black", font=heading_font)

    top = sorted(context["outlet_counts"].items(), key=lambda x: (-x[1], x[0]))[:12]
    max_count = max([v for _, v in top], default=1)
    chart_x = 250
    bar_max = 950
    y = 365

    for outlet, count in top:
        draw.text((60, y + 5), str(outlet), fill="black", font=normal_font)
        bar_width = int((count / max_count) * bar_max)
        draw.rectangle((chart_x, y, chart_x + max(bar_width, 3), y + 28), outline="black", width=2)
        draw.text((chart_x + max(bar_width, 3) + 10, y + 3), str(count), fill="black", font=normal_font)
        y += 45

    if not top:
        draw.text((60, 380), "No records for this report.", fill="black", font=normal_font)

    footer = f"Generated {display_datetime(context['generated_at'])}"
    draw.text((60, height - 55), footer, fill="black", font=small_font)

    image.save(path, "PNG")
    return str(path)


def set_report_query(user_id, chat_id, rows, start_date=None, end_date=None, outlet=None, title="SGSB BAGGING REPORT"):
    context = build_report_context(rows, start_date, end_date, outlet, title)
    report_sessions[user_id] = {
        "step": "format",
        "context": context
    }
    send_message(
        chat_id,
        "✅ <b>REPORT DATA READY</b>\n\n"
        f"📅 {safe_html(report_period_text(context))}\n"
        f"🏢 {safe_html(outlet or 'Semua Outlet')}\n"
        f"📦 Records: <b>{len(context['active'])}</b>\n\n"
        "Pilih format report:",
        report_format_keyboard()
    )


def handle_report_text(chat_id, user_id, text):
    session = report_sessions.get(user_id)
    if not session:
        return False

    step = session.get("step")

    if step == "passcode":
        if ADMIN_PASSCODE and secrets.compare_digest(text, ADMIN_PASSCODE):
            admin_sessions[user_id] = {"authenticated": True, "step": "menu"}
            report_sessions[user_id] = {"step": "menu"}
            send_message(chat_id, "✅ Access dibenarkan.\n\nPilih jenis report:", report_menu_keyboard())
        else:
            report_sessions.pop(user_id, None)
            send_message(chat_id, "❌ Passcode salah.")
        return True

    if step == "range_input":
        match = re.match(r"^\s*(.+?)\s*(?:-|hingga|to)\s*(.+?)\s*$", text, re.I)
        if not match:
            send_message(chat_id, "⚠️ Format: <code>01/09/2026 - 20/09/2026</code>")
            return True

        start = parse_date_input(match.group(1))
        end = parse_date_input(match.group(2))
        if not start or not end or start > end:
            send_message(chat_id, "⚠️ Tarikh tidak sah. Contoh: <code>01/09/2026 - 20/09/2026</code>")
            return True

        rows = load_records_range(start, end)
        set_report_query(user_id, chat_id, rows, start, end)
        return True

    if step == "outlet_range_input":
        if "|" not in text:
            send_message(chat_id, "⚠️ Format: <code>SBH307 | 01/09/2026 - 20/09/2026</code>")
            return True

        outlet_part, range_part = [x.strip() for x in text.split("|", 1)]
        outlet = normalize_outlet(outlet_part)
        match = re.match(r"^\s*(.+?)\s*(?:-|hingga|to)\s*(.+?)\s*$", range_part, re.I)
        if not outlet or not match:
            send_message(chat_id, "⚠️ Format: <code>SBH307 | 01/09/2026 - 20/09/2026</code>")
            return True

        start = parse_date_input(match.group(1))
        end = parse_date_input(match.group(2))
        if not start or not end or start > end:
            send_message(chat_id, "⚠️ Tarikh tidak sah.")
            return True

        rows = load_records_range(start, end, outlet=outlet)
        set_report_query(user_id, chat_id, rows, start, end, outlet=outlet)
        return True

    return False


def deliver_report_format(chat_id, user_id, fmt):
    session = report_sessions.get(user_id)
    if not session or session.get("step") != "format":
        send_message(chat_id, "⚠️ Report session tamat. Taip /report semula.")
        return

    context = session.get("context") or {}

    if fmt == "summary":
        text = report_summary_text(context, html_mode=True)
        if len(text) <= 3900:
            send_message(chat_id, text, report_format_keyboard())
        else:
            send_long_plain(chat_id, report_summary_text(context, html_mode=False))
            send_message(chat_id, "Pilih format lain jika diperlukan:", report_format_keyboard())
        return

    if fmt == "detail":
        send_long_plain(chat_id, report_detail_text(context))
        send_message(chat_id, "✅ Full detail selesai dihantar.", report_format_keyboard())
        return

    if fmt == "archive":
        text = report_summary_text(context, html_mode=True)
        if len(text) > 3900:
            send_message(chat_id, "⚠️ Summary terlalu panjang untuk satu post Archive. Gunakan PDF/Excel/TXT.")
            return
        result = send_message(ARCHIVE_CHAT_ID, text)
        if result.get("ok"):
            send_message(chat_id, "✅ Summary telah dipost ke Archive Channel.", report_format_keyboard())
        else:
            send_message(chat_id, f"❌ Gagal post ke Archive: <code>{safe_html(result.get('description'))}</code>")
        return

    generators = {
        "txt": generate_txt_report,
        "excel": generate_excel_report,
        "pdf": generate_pdf_report,
        "chart": generate_chart_report,
    }

    generator = generators.get(fmt)
    if not generator:
        return

    send_message(chat_id, f"⏳ Menjana <b>{safe_html(fmt.upper())}</b> report...")
    filepath = None
    try:
        filepath = generator(context)
        caption = f"SGSB Bagging Report | {report_period_text(context)}"
        if fmt == "chart":
            result = send_photo_file(chat_id, filepath, caption)
        else:
            result = send_document_file(chat_id, filepath, caption)

        if not result.get("ok"):
            send_message(chat_id, f"❌ Gagal menghantar report: <code>{safe_html(result.get('description'))}</code>")
        else:
            send_message(chat_id, "✅ Report siap.", report_format_keyboard())
    except ImportError as e:
        send_message(
            chat_id,
            "❌ Library report belum terpasang pada Render.\n"
            "Pastikan requirements.txt V5 sudah digunakan.\n"
            "Sila cuba semula atau hubungi admin."
        )
    except Exception as e:
        print("Report generation error:", type(e).__name__, flush=True)
        send_message(chat_id, "❌ Report gagal dijana. Sila cuba semula atau hubungi admin.")
    finally:
        if filepath:
            try:
                os.remove(filepath)
            except Exception:
                pass

# ============================================================
# MESSAGE HANDLER
# ============================================================

def handle_message(message):
    chat_id = message["chat"]["id"]
    user = message.get("from", {})
    user_id = user.get("id")

    if not user_id:
        return

    cleanup_old_runtime_records()
    text = message.get("text", "").strip()

    # Commands
    if text == "/admin":
        begin_admin(chat_id, user_id)
        return

    if text == "/report":
        begin_report(chat_id, user_id)
        return

    if text == "/late":
        start_builder(chat_id, user_id, mode="LATE")
        return

    if text in ("/start", "/new"):
        start_builder(chat_id, user_id, mode="NORMAL")
        return

    # Report text input before admin text input.
    if handle_report_text(chat_id, user_id, text):
        return

    if handle_admin_text(chat_id, user_id, text):
        return

    # Remark input.
    mode = input_modes.get(user_id)
    if mode and mode.get("type") == "remark":
        rid = mode.get("rid")
        record = records.get(rid) or get_record_persistent(rid)

        if not record:
            input_modes.pop(user_id, None)
            send_message(chat_id, "⚠️ Record tidak lagi tersedia.")
            return

        if not text:
            send_message(chat_id, "⚠️ Taip remark dalam bentuk teks.")
            return

        record["remark"] = text[:1000]
        input_modes.pop(user_id, None)
        with state_lock:
            records[rid] = record
        ok, error = upsert_record_persistent(record)
        update_record_card(record)
        if ok:
            send_message(chat_id, "✅ Remark disimpan.")
        else:
            send_message(chat_id, f"⚠️ Remark di memory tetapi storage gagal: <code>{safe_html(error)}</code>")
        return

    # Media handling.
    if "photo" in message or "video" in message or "document" in message:
        record = prepare_media(message, user_id, chat_id)
        if not record:
            return

        if record.get("status") in ("COMPLETE", "LATE_COMPLETE"):
            send_message(chat_id, "ℹ️ Record ini sudah COMPLETE dan berada dalam Archive.")
            return

        if record.get("status") == "CANCELLED":
            send_message(chat_id, "ℹ️ Record ini sudah dibatalkan.")
            return

        if "photo" in message:
            if record.get("photo"):
                if record["photo"].get("message_id") == message["message_id"]:
                    maybe_archive_record(record)
                    return
                send_message(chat_id, "⚠️ Record ini sudah mempunyai 1 gambar.")
                return

            photo = message["photo"][-1]
            record["photo"] = {
                "file_id": photo["file_id"],
                "message_id": message["message_id"]
            }
            if record.get("photo") and record.get("video") and not record.get("completed_at"):
                record["completed_at"] = now_my().isoformat()
            with state_lock:
                records[record["rid"]] = record
            send_message(chat_id, f"📷 Gambar diterima untuk <b>{safe_html(record['seal'])}</b>.")
            maybe_archive_record(record)
            return

        if "video" in message:
            if record.get("video"):
                if record["video"].get("message_id") == message["message_id"]:
                    maybe_archive_record(record)
                    return
                send_message(chat_id, "⚠️ Record ini sudah mempunyai 1 video.")
                return

            video = message["video"]
            record["video"] = {
                "file_id": video["file_id"],
                "message_id": message["message_id"]
            }
            if record.get("photo") and record.get("video") and not record.get("completed_at"):
                record["completed_at"] = now_my().isoformat()
            with state_lock:
                records[record["rid"]] = record
            send_message(chat_id, f"🎥 Video diterima untuk <b>{safe_html(record['seal'])}</b>.")
            maybe_archive_record(record)
            return

        if "document" in message:
            doc = message["document"]
            mime = doc.get("mime_type", "")
            if mime.startswith("video/"):
                send_message(
                    chat_id,
                    "⚠️ Untuk memastikan gambar + video menjadi <b>satu album</b> di Archive, "
                    "hantar video sebagai <b>Video</b>, bukan File/Document."
                )
                return

            send_message(chat_id, "⚠️ Fail ini tidak digunakan untuk Bagging Record.")
            return

    # Builder text input.
    builder = builders.get(user_id)

    if builder and builder.get("step") == "seal":
        if not text:
            send_message(chat_id, "⚠️ Masukkan Seal Bag Number.")
            return

        seal = normalize_seal(text)
        if not seal:
            send_message(chat_id, "⚠️ Seal Bag Number tidak sah.")
            return

        builder["seal"] = seal

        if builder.get("mode") == "NORMAL":
            complete_normal_builder(chat_id, user_id, builder, seal)
            return

        builder["step"] = "late_date"
        send_message(
            chat_id,
            "📅 <b>Tarikh sebenar Bagging</b>\n\n"
            "Taip tarikh, contoh <code>18/09/2026</code>, atau pilih:",
            late_date_keyboard()
        )
        return

    if builder and builder.get("step") == "late_date":
        selected = parse_date_input(text)
        if not selected:
            send_message(chat_id, "⚠️ Format tarikh: <code>DD/MM/YYYY</code>", late_date_keyboard())
            return
        if selected > now_my().date():
            send_message(chat_id, "⚠️ Bagging Date tidak boleh tarikh masa depan.")
            return
        handle_late_date(chat_id, user_id, selected)
        return

    if builder and builder.get("step") == "late_reason":
        reason = "-" if text == "/skip" else text
        if not reason:
            send_message(chat_id, "⚠️ Taip sebab atau /skip.")
            return
        finalize_late_builder(chat_id, user_id, reason)
        return

    if not builder:
        send_message(
            chat_id,
            "👋 <b>SGSB Bagging Bot V5.1.1</b>\n\n"
            "Bagging baru: /start\n"
            "Late / missing media: /late\n"
            "Report: /report\n"
            "Admin: /admin"
        )

# ============================================================
# CALLBACK HANDLER
# ============================================================

def handle_callback(callback):
    callback_id = callback["id"]
    user = callback["from"]
    user_id = user["id"]
    message = callback.get("message", {})
    chat_id = message.get("chat", {}).get("id")
    data = callback.get("data", "")

    if not callback.get("_btm_answered"):
        answer_callback(callback_id)

    if data.startswith("media:"):
        parts = data.split(":")
        if len(parts) == 3:
            choose_pending_media(callback, parts[1], parts[2], user_id, chat_id)
        return

    if data == "noop":
        return

    if data == "new":
        start_builder(chat_id, user_id, mode="NORMAL")
        return

    # Normal outlet selection.
    if data.startswith("outlet:"):
        outlet = data.split(":", 1)[1]
        with state_lock:
            valid = outlet in OUTLETS
        if not valid:
            send_message(chat_id, "⚠️ Outlet tidak sah / sudah dipadam.")
            return

        builders[user_id] = {
            "mode": "NORMAL",
            "step": "seal",
            "outlet": outlet,
            "seal": None,
            "bagging_date": now_my().date().isoformat(),
            "late_reason": ""
        }
        send_message(chat_id, f"🏢 Outlet: <b>{safe_html(outlet)}</b>\n\nMasukkan <b>Seal Bag Number</b>.")
        return

    # Late outlet selection.
    if data.startswith("late_outlet:"):
        outlet = data.split(":", 1)[1]
        with state_lock:
            valid = outlet in OUTLETS
        if not valid:
            send_message(chat_id, "⚠️ Outlet tidak sah / sudah dipadam.")
            return

        builders[user_id] = {
            "mode": "LATE",
            "step": "seal",
            "outlet": outlet,
            "seal": None,
            "bagging_date": None,
            "late_reason": ""
        }
        send_message(chat_id, f"🏢 Outlet: <b>{safe_html(outlet)}</b>\n\nMasukkan <b>Seal Bag Number</b> yang terlupa/lambat.")
        return

    if data.startswith("late_date:"):
        builder = builders.get(user_id)
        if not builder or builder.get("mode") != "LATE" or builder.get("step") != "late_date":
            send_message(chat_id, "⚠️ Late session sudah tamat. Taip /late semula.")
            return
        choice = data.split(":", 1)[1]
        selected = now_my().date() if choice == "today" else now_my().date() - timedelta(days=1)
        handle_late_date(chat_id, user_id, selected)
        return

    # Record callbacks.
    if data.startswith("remark:"):
        rid = data.split(":", 1)[1]
        record = records.get(rid) or get_record_persistent(rid)
        if not record:
            send_message(chat_id, "⚠️ Record tidak dijumpai.")
            return

        if record.get("owner_id") != user_id and not admin_authenticated(user_id):
            send_message(chat_id, "⚠️ Rekod ini milik pengguna lain.")
            return
        input_modes[user_id] = {"type": "remark", "rid": rid}
        send_message(chat_id, f"📝 Taip remark untuk Seal <b>{safe_html(record.get('seal'))}</b>.")
        return

    if data.startswith("retry:"):
        rid = data.split(":", 1)[1]
        record = records.get(rid) or get_record_persistent(rid)
        if not record:
            send_message(chat_id, "⚠️ Record tidak dijumpai.")
            return

        if record.get("owner_id") != user_id and not admin_authenticated(user_id):
            send_message(chat_id, "⚠️ Rekod ini milik pengguna lain.")
            return
        record["owner_id"] = user_id
        record["chat_id"] = chat_id
        record["archiving"] = False
        with state_lock:
            records[rid] = record
        ok, error = archive_record(record)
        if not ok and error != "Archive sedang diproses":
            send_message(chat_id, f"⚠️ Archive masih gagal.\n<code>{safe_html(error)}</code>")
        return

    if data.startswith("cancel_record:"):
        rid = data.split(":", 1)[1]
        record = records.get(rid) or get_record_persistent(rid)
        if not record:
            send_message(chat_id, "⚠️ Record tidak dijumpai.")
            return

        if record.get("status") in ("COMPLETE", "LATE_COMPLETE"):
            send_message(chat_id, "⚠️ Record yang sudah di-Archive tidak boleh dibatalkan.")
            return

        record["status"] = "CANCELLED"
        record["cancelled_at"] = now_my().isoformat()
        if record.get("owner_id") != user_id and not admin_authenticated(user_id):
            send_message(chat_id, "⚠️ Rekod ini milik pengguna lain.")
            return
        record["owner_id"] = user_id
        record["chat_id"] = chat_id
        upsert_record_persistent(record)

        with state_lock:
            records[record["rid"]] = record

        edit_message(
            chat_id,
            record.get("card_message_id"),
            "❌ <b>BAGGING RECORD DIBATALKAN</b>\n\n"
            f"🏢 {safe_html(record.get('outlet'))}\n"
            f"🔒 {safe_html(record.get('seal'))}\n"
            f"🆔 <code>{safe_html(record_display_id(record))}</code>",
            {"inline_keyboard": [[{"text": "➕ NEW BAGGING", "callback_data": "new"}]]}
        )
        return

    # Report callbacks.
    if data.startswith("report:") and not admin_authenticated(user_id):
        begin_report(chat_id, user_id)
        return
    if data == "report:menu":
        if not admin_authenticated(user_id):
            begin_report(chat_id, user_id)
        else:
            report_sessions[user_id] = {"step": "menu"}
            send_message(chat_id, "📊 <b>SGSB BAGGING REPORT</b>\n\nPilih jenis report:", report_menu_keyboard())
        return

    if data == "report:close":
        report_sessions.pop(user_id, None)
        send_message(chat_id, "✅ Report menu ditutup.")
        return

    if data == "report:q:today":
        d = now_my().date()
        rows = load_records_range(d, d)
        set_report_query(user_id, chat_id, rows, d, d)
        return

    if data == "report:q:yesterday":
        d = now_my().date() - timedelta(days=1)
        rows = load_records_range(d, d)
        set_report_query(user_id, chat_id, rows, d, d)
        return

    if data == "report:q:range":
        report_sessions[user_id] = {"step": "range_input"}
        send_message(chat_id, "🗓 Taip Date Range:\n<code>01/09/2026 - 20/09/2026</code>")
        return

    if data == "report:q:outlet":
        report_sessions[user_id] = {"step": "outlet_range_input"}
        send_message(chat_id, "🏢 Taip Outlet + Date Range:\n<code>SBH307 | 01/09/2026 - 20/09/2026</code>")
        return

    if data == "report:q:incomplete":
        rows = [
            r for r in load_all_records()
            if r.get("status") not in ("COMPLETE", "LATE_COMPLETE", "CANCELLED")
            or not r.get("photo")
            or not r.get("video")
        ]
        set_report_query(user_id, chat_id, rows, title="SGSB INCOMPLETE BAGGING REPORT")
        return

    if data == "report:q:late":
        rows = [
            r for r in load_all_records()
            if r.get("entry_mode") == "LATE" or r.get("status") == "LATE_COMPLETE"
        ]
        set_report_query(user_id, chat_id, rows, title="SGSB LATE BAGGING REPORT")
        return

    if data.startswith("report:fmt:"):
        fmt = data.split(":", 2)[2]
        deliver_report_format(chat_id, user_id, fmt)
        return

    # Admin callbacks.
    if data.startswith("admin:"):
        adm = admin_sessions.get(user_id)
        if not adm or not adm.get("authenticated"):
            send_message(chat_id, "🔐 Sesi admin tamat. Taip /admin semula.")
            return

        if data == "admin:menu":
            show_admin_menu(chat_id, user_id)
            return

        if data == "admin:list":
            with state_lock:
                outlets = sorted(OUTLETS)
            outlet_text = "\n".join(f"• {safe_html(x)}" for x in outlets) or "Tiada outlet."
            storage_text = "✅ Permanent" if OUTLET_STORAGE_OK else "⚠️ Tidak permanent"
            send_message(
                chat_id,
                "📋 <b>SENARAI OUTLET</b>\n\n"
                f"{outlet_text}\n\nTotal: {len(outlets)}\nStorage: {storage_text}",
                admin_keyboard()
            )
            return

        if data == "admin:add":
            adm["step"] = "add"
            send_message(chat_id, "➕ Taip <b>Outlet Code</b> baru.\nContoh: SBH999")
            return

        if data == "admin:delete":
            send_message(chat_id, "🗑 <b>Pilih outlet untuk dipadam:</b>", delete_outlet_keyboard())
            return

        if data.startswith("admin:del:"):
            outlet = data.split(":", 2)[2]
            with state_lock:
                current = list(OUTLETS)

            if outlet not in current:
                send_message(chat_id, "ℹ️ Outlet tidak dijumpai.", admin_keyboard())
                return

            updated = [x for x in current if x != outlet]
            ok, error = save_outlets_to_supabase(updated)
            if not ok:
                send_message(
                    chat_id,
                    "❌ <b>OUTLET TIDAK DIPADAM</b>\n"
                    f"<code>{safe_html(error)}</code>",
                    admin_keyboard()
                )
                return

            with state_lock:
                OUTLETS[:] = sorted(updated)

            send_message(chat_id, f"✅ <b>{safe_html(outlet)}</b> dipadam secara permanent.", admin_keyboard())
            return

        if data == "admin:sync":
            ok = load_outlets_from_supabase()
            if ok:
                send_message(chat_id, f"✅ Outlet Master synced.\nTotal: {len(OUTLETS)}", admin_keyboard())
            else:
                send_message(chat_id, f"❌ Sync gagal.\n<code>{safe_html(OUTLET_STORAGE_ERROR)}</code>", admin_keyboard())
            return

        if data == "admin:logout":
            admin_sessions.pop(user_id, None)
            report_sessions.pop(user_id, None)
            send_message(chat_id, "🚪 Admin logout.")
            return

# ============================================================
# WEB ROUTES
# ============================================================

@app.route("/", methods=["GET"])
def home():
    return jsonify({
        "status": "ok",
        "service": "SGSB Bagging Bot V5.1.1",
        "outlets": len(OUTLETS),
        "outlet_storage": "supabase" if OUTLET_STORAGE_OK else "unavailable",
        "metadata_storage": "supabase",
        "features": [
            "parallel-record-cards",
            "persistent-record-log",
            "late-upload",
            "detailed-reports",
            "pdf",
            "excel",
            "txt",
            "chart"
        ]
    })


@app.route("/health", methods=["GET"])
def health():
    connected=load_outlets_from_supabase()
    return jsonify({'status':'healthy' if connected else 'degraded','version':'5.1.1','storage':'supabase','database_connected':connected}), 200 if connected else 503


@app.route("/webhook", methods=["POST"])
def webhook():
    if not WEBHOOK_SECRET or not secrets.compare_digest(request.headers.get("X-Telegram-Bot-Api-Secret-Token", ""), WEBHOOK_SECRET):
        return jsonify({"ok": False}), 403
    update = request.get_json(silent=True) or {}
    update_id=update.get('update_id')
    if not isinstance(update_id,int): return jsonify({'ok':False}),400
    with update_lock(update):
        callback=update.get('callback_query') or {}
        if callback.get('id'):
            acknowledged=answer_callback(callback['id'])
            if acknowledged.get('ok'): callback['_btm_answered']=True
        token=secrets.token_hex(16)
        claimed=False
        try:
            claim=db.rpc('sgsb_claim_update',{'p_id':update_id,'p_token':token})
            if claim=='done': return jsonify({'ok':True,'duplicate':True})
            if claim!='claimed': return jsonify({'ok':False,'retry':True}),503
            claimed=True
            if not ensure_outlets_available(): raise db.StorageError(db.FRIENDLY)
            if report_menu.handle(update): pass
            elif fleet_bot.handle(update): pass
            elif "message" in update: handle_message(update["message"])
            elif "callback_query" in update: handle_callback(update["callback_query"])
            db.rpc('sgsb_finish_update',{'p_id':update_id,'p_token':token,'p_done':True})
            return jsonify({'ok':True})
        except Exception as e:
            print('Webhook processing failed:',type(e).__name__,flush=True)
            if claimed:
                try: db.rpc('sgsb_finish_update',{'p_id':update_id,'p_token':token,'p_done':False})
                except db.StorageError: pass
            msg=update.get('message') or update.get('callback_query',{}).get('message',{})
            chat_id=msg.get('chat',{}).get('id')
            if chat_id: send_message(chat_id,'⚠️ '+db.FRIENDLY)
            return jsonify({'ok':False,'retry':True}),503

# ============================================================
# STARTUP
# ============================================================

fleet_bot = install_fleet(app, db, telegram, admin_sessions)

def open_bagging_from_menu(chat_id,user_id):
    builder=builders.get(user_id)
    if not builder:
        start_builder(chat_id,user_id,mode='NORMAL');return
    step=builder.get('step')
    if step=='outlet':
        send_message(chat_id,'Continue your existing Video Bagging report. Choose outlet:',outlet_keyboard('late_outlet' if builder.get('mode')=='LATE' else 'outlet'))
    elif step=='seal':
        send_message(chat_id,'Continue Video Bagging for <b>'+safe_html(builder.get('outlet'))+'</b>. Enter the Seal Bag Number.')
    elif step=='late_date':
        send_message(chat_id,'Continue your existing backdated report. Enter its original date (DD/MM/YYYY).',late_date_keyboard())
    elif step=='late_reason':
        send_message(chat_id,'Continue your existing backdated report. Enter the reason, or /skip.')
    else:
        send_message(chat_id,'Your existing Bagging input is kept. Continue from its original message; /menu does not reset it.')

report_menu = install_report_menu(db,telegram,fleet_bot,open_bagging_from_menu)
load_outlets_from_supabase()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=port)
